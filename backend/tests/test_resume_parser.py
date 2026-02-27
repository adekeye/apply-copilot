import io

from docx import Document

from app.services.resume_parser import parse_resume


def test_parse_resume_from_text_sections():
    raw = """
    Jane Engineer
    jane@example.com
    +1 (415) 555-0101
    https://linkedin.com/in/jane

    Skills:
    Python, FastAPI, AWS, SQL, React

    Experience:
    Senior Engineer, Data Systems Inc (2018 - Present)

    Projects:
    - Built hiring copilot workflows

    Education:
    University of Example, B.S. Computer Science

    Authorized to work in the United States. Open to remote and Seattle roles.
    """
    profile = parse_resume("resume.txt", raw.encode("utf-8"))
    assert profile.contact["email"] == "jane@example.com"
    assert "python" in profile.skills
    assert "aws" in profile.skills
    assert profile.years_experience >= 6
    assert "Remote" in profile.location_preferences
    assert profile.work_auth == "Authorized"


def test_parse_resume_from_docx():
    doc = Document()
    doc.add_paragraph("Alex Candidate")
    doc.add_paragraph("alex@example.com")
    doc.add_paragraph("Skills:")
    doc.add_paragraph("TypeScript, Next.js, Docker, Kubernetes")
    doc.add_paragraph("Experience:")
    doc.add_paragraph("Staff Engineer, Platform Technologies LLC (2017 - 2024)")

    buf = io.BytesIO()
    doc.save(buf)

    profile = parse_resume("resume.docx", buf.getvalue())
    assert profile.contact["email"] == "alex@example.com"
    assert "typescript" in profile.skills
    assert "next.js" in profile.skills
    assert any("Technologies LLC" in employer for employer in profile.employers)
    assert profile.years_experience >= 7
