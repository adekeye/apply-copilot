import io
import re
from pathlib import Path
from typing import Iterable

from docx import Document
from pypdf import PdfReader

from app.schemas import ResumeStructured


SKILL_ALIASES = {
    "python": {"python"},
    "sql": {"sql", "mysql", "postgres", "postgresql"},
    "typescript": {"typescript", "ts"},
    "javascript": {"javascript", "js"},
    "react": {"react", "reactjs", "react.js"},
    "next.js": {"next.js", "nextjs"},
    "aws": {"aws", "amazon web services"},
    "docker": {"docker"},
    "kubernetes": {"kubernetes", "k8s"},
    "fastapi": {"fastapi"},
    "node": {"node", "nodejs", "node.js"},
    "postgresql": {"postgresql", "postgres"},
    "ml": {"ml", "machine learning"},
    "nlp": {"nlp", "natural language processing"},
    "java": {"java"},
    "go": {"go", "golang"},
    "spark": {"spark", "apache spark"},
}

SECTION_HEADERS = (
    "experience",
    "professional experience",
    "work experience",
    "education",
    "skills",
    "projects",
    "certifications",
    "summary",
)


def _extract_text(filename: str, content: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        pages: list[str] = []
        for page in reader.pages:
            page_text = page.extract_text(extraction_mode="layout") or page.extract_text() or ""
            pages.append(page_text)
        return "\n".join(pages)
    if ext == ".docx":
        doc = Document(io.BytesIO(content))
        lines: list[str] = []
        lines.extend(p.text for p in doc.paragraphs if p.text and p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    return content.decode("utf-8", errors="ignore")


def _normalize_text(text: str) -> str:
    normalized = text.replace("\xa0", " ")
    normalized = re.sub(r"[ \t]+", " ", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _clean_link(raw: str) -> str:
    return raw.rstrip(".,);]")


def _find_skills(text: str) -> list[str]:
    lowered = text.lower()
    found: list[str] = []
    for canonical, aliases in SKILL_ALIASES.items():
        if any(re.search(rf"\b{re.escape(alias)}\b", lowered) for alias in aliases):
            found.append(canonical)
    return sorted(found)


def _section_lines(lines: list[str], section_name: str) -> list[str]:
    start = None
    section = section_name.lower()
    for idx, line in enumerate(lines):
        if line.lower().strip(":") == section:
            start = idx + 1
            break
    if start is None:
        return []

    out: list[str] = []
    for line in lines[start:]:
        normalized = line.lower().strip(":")
        if normalized in SECTION_HEADERS:
            break
        out.append(line)
    return out


def _iter_bullets(lines: Iterable[str]) -> list[str]:
    bullets: list[str] = []
    for line in lines:
        if re.match(r"^\s*[-*•]\s+", line):
            bullets.append(re.sub(r"^\s*[-*•]\s+", "", line).strip())
    return bullets


def _extract_years_experience(text: str) -> int:
    lowered = text.lower()
    direct = [int(v) for v in re.findall(r"(\d{1,2})\+?\s*(?:years|yrs)", lowered)]
    if direct:
        return max(direct)

    # Fallback from date ranges such as "2018 - 2024"
    ranges = re.findall(r"(20\d{2})\s*[-–]\s*(20\d{2}|present|current)", lowered)
    best = 0
    for start, end in ranges:
        start_i = int(start)
        end_i = 2026 if end in {"present", "current"} else int(end)
        if end_i >= start_i:
            best = max(best, end_i - start_i)
    return best


def _extract_employers(experience_lines: list[str]) -> list[str]:
    employers: list[str] = []
    for line in experience_lines:
        if re.search(r"\b(inc|llc|corp|company|technologies|systems|labs|group)\b", line.lower()):
            employers.append(line.strip())
    return employers[:8]


def parse_resume(filename: str, content: bytes) -> ResumeStructured:
    text = _normalize_text(_extract_text(filename, content))
    lowered = text.lower()
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    email_match = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone_match = re.search(r"(\+?\d[\d\s().-]{8,}\d)", text[:1200])
    links = [_clean_link(l) for l in re.findall(r"https?://\S+|www\.\S+", text)]

    skills_section = _section_lines(lines, "skills")
    skills = _find_skills("\n".join(skills_section) if skills_section else text)
    years_experience = _extract_years_experience(text)

    project_lines = _section_lines(lines, "projects")
    projects = (_iter_bullets(project_lines) or project_lines)[:8]

    education_lines = _section_lines(lines, "education")
    if not education_lines:
        education_lines = [l for l in lines if any(k in l.lower() for k in ["university", "college", "bachelor", "master", "phd", "b.s.", "m.s."])]
    education = education_lines[:6]

    experience_lines = _section_lines(lines, "experience") or _section_lines(lines, "professional experience") or _section_lines(lines, "work experience")
    employers = _extract_employers(experience_lines if experience_lines else lines)

    keywords = sorted(set(re.findall(r"[A-Za-z][A-Za-z0-9+.#-]{2,}", text)))[:50]

    location_prefs: list[str] = []
    for loc in ["remote", "new york", "san francisco", "seattle", "austin", "united states", "boston", "chicago"]:
        if re.search(rf"\b{re.escape(loc)}\b", lowered):
            location_prefs.append(loc.title())

    work_auth = "Unknown"
    if re.search(r"\b(authorized to work|us citizen|green card)\b", lowered):
        work_auth = "Authorized"
    if re.search(r"\b(require|need).{0,20}sponsorship\b", lowered):
        work_auth = "Needs sponsorship"

    return ResumeStructured(
        contact={"email": email_match.group(0) if email_match else None, "phone": phone_match.group(0) if phone_match else None},
        skills=skills,
        years_experience=years_experience,
        projects=projects,
        employers=employers,
        education=education,
        keywords=keywords,
        location_preferences=location_prefs,
        work_auth=work_auth,
        links=links,
        raw_text=text,
    )
